import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# Load cleaned CSV (skip the junk header row)
df = pd.read_csv("output.csv", skiprows=1)

# Clean column names
df.columns = [c.replace("\n", " ").strip().upper() for c in df.columns]

# Convert CET PERCENTILE to numeric (force errors to NaN)
df["CET PERCENTILE"] = pd.to_numeric(df["CET PERCENTILE"], errors="coerce")

# Fill missing values with 0 (or use dropna() if you want to ignore them completely)
df["CET PERCENTILE"] = df["CET PERCENTILE"].fillna(0)

# Sort by CET Percentile
df_sorted = df.sort_values(by="CET PERCENTILE", ascending=False)

# Add Rank column (works now even with NaN handled)
df_sorted["RANK"] = df_sorted["CET PERCENTILE"].rank(ascending=False, method="dense").astype(int)

# Save to Excel
excel_file = "ranked_output.xlsx"
df_sorted.to_excel(excel_file, index=False, sheet_name="Ranking")

# Open with openpyxl for formatting
wb = load_workbook(excel_file)
ws = wb["Ranking"]

# Styles
header_font = Font(bold=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

# Format header
for cell in ws[1]:
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

# Auto column width + borders
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value is not None:
            max_len = max(max_len, len(str(cell.value)))
        cell.border = thin_border
    ws.column_dimensions[col_letter].width = max_len + 2

# Freeze header row
ws.freeze_panes = "A2"

wb.save(excel_file)
print("✅ Ranked & formatted Excel saved as", excel_file)
